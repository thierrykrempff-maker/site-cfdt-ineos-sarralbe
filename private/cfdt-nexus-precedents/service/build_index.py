#!/usr/bin/env python
"""Construit (ou complète) l'index texte PRIVÉ des PV du CE/CSE.

Lecture seule sur la source ; écriture uniquement dans l'index (hors dépôt). Reprenable : un fichier déjà indexé est ignoré.

Formats : PDF (texte natif via pdftotext, sinon OCR français via pdftoppm + Tesseract), DOCX (python-docx, avec
récupération partielle si le fichier est abîmé), PPTX, DOC (antiword). Les .ppt anciens ne sont pas lus.

Configuration (variables d'environnement) :
  NEXUS_PV_SOURCE   dossier des PV                  (défaut : D:\\pv cse\\Réunions de CE)
  NEXUS_PV_INDEX    dossier de l'index privé        (défaut : ~/.claude/nexus-private/pv_index)
  TESSERACT_EXE     chemin de tesseract.exe         (défaut : C:\\Program Files\\Tesseract-OCR\\tesseract.exe)
  YEAR_MIN          n'indexe que les dossiers « PV du CE/CSE <année> » >= cette année

Prérequis : poppler (pdftotext, pdftoppm) dans le PATH, Tesseract avec le pack « fra », python-docx, antiword.

Usage : python -I build_index.py [nb_threads]     puis     python -I precedents.py --rebuild --q "..."
"""
import os, sys, re, json, html, zlib, hashlib, zipfile, tempfile, shutil, subprocess, threading
from concurrent.futures import ThreadPoolExecutor

SRC = os.environ.get("NEXUS_PV_SOURCE") or r"D:\pv cse\Réunions de CE"
OUT = os.environ.get("NEXUS_PV_INDEX") or os.path.join(os.path.expanduser("~"), ".claude", "nexus-private", "pv_index")
TXT = os.path.join(OUT, "txt")
LOG = os.path.join(OUT, "index.jsonl")
TESS = os.environ.get("TESSERACT_EXE") or r"C:\Program Files\Tesseract-OCR\tesseract.exe"
EXTS = {".pdf", ".docx", ".pptx", ".doc"}
NOWIN = getattr(subprocess, "CREATE_NO_WINDOW", 0)
LOCK = threading.Lock()
os.environ["OMP_THREAD_LIMIT"] = "1"  # Tesseract : un cœur par tâche, sinon sur-souscription


def run(cmd, timeout):
    return subprocess.run(cmd, capture_output=True, timeout=timeout, creationflags=NOWIN)


def pdf_native(p):
    return run(["pdftotext", "-enc", "UTF-8", "-layout", p, "-"], 120).stdout.decode("utf-8", "replace")


def pdf_ocr(p):
    tmp = tempfile.mkdtemp(dir=OUT)
    try:
        run(["pdftoppm", "-r", "200", "-gray", "-png", p, os.path.join(tmp, "pg")], 600)
        parts = []
        for i, f in enumerate(sorted(x for x in os.listdir(tmp) if x.endswith(".png")), 1):
            r = run([TESS, os.path.join(tmp, f), "stdout", "-l", "fra", "--psm", "3"], 300)
            parts.append(f"\n--- page {i} ---\n" + r.stdout.decode("utf-8", "replace"))
        return "".join(parts)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def docx_text(p):
    import docx
    d = docx.Document(p)
    out = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


def docx_salvage(p):
    """Récupération partielle d'un .docx dont le CRC est invalide : décompression brute de document.xml."""
    z = zipfile.ZipFile(p)
    zi = z.getinfo("word/document.xml")
    with open(p, "rb") as f:
        f.seek(zi.header_offset)
        h = f.read(30)
        f.seek(zi.header_offset + 30 + int.from_bytes(h[26:28], "little") + int.from_bytes(h[28:30], "little"))
        data = f.read(zi.compress_size)
    d, out = zlib.decompressobj(-15), b""
    for i in range(0, len(data), 512):
        try:
            out += d.decompress(data[i:i + 512])
        except Exception:
            break
    s = re.sub(r"</w:p>", "\n", out.decode("utf-8", "replace"))
    return html.unescape(re.sub(r"\s*\n\s*", "\n", re.sub(r"<[^>]+>", "", s)))


def pptx_text(p):
    out = []
    with zipfile.ZipFile(p) as z:
        names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for n in names:
            x = z.read(n).decode("utf-8", "replace")
            out.append(f"--- {n} ---\n" + " ".join(re.findall(r"<a:t>([^<]*)</a:t>", x)))
    return "\n".join(out)


def doc_text(p):
    return run(["antiword", "-m", "UTF-8.txt", p], 60).stdout.decode("utf-8", "replace")


def extract(p):
    ext = os.path.splitext(p)[1].lower()
    if ext == ".pdf":
        t = pdf_native(p)
        return (t, "native") if len(re.sub(r"\s", "", t)) >= 100 else (pdf_ocr(p), "ocr")
    if ext == ".docx":
        try:
            return docx_text(p), "docx"
        except Exception:
            return docx_salvage(p), "docx-salvage"
    if ext == ".pptx":
        return pptx_text(p), "pptx"
    return doc_text(p), "doc"


def process(p):
    h = hashlib.md5(p.encode("utf-8")).hexdigest()
    dst = os.path.join(TXT, h + ".txt")
    if os.path.exists(dst) or os.path.basename(p).startswith("~$"):  # fichiers verrou Office ignorés
        return
    text, method, err = "", "", ""
    try:
        text, method = extract(p)
    except Exception as e:
        err, method = repr(e)[:200], "error"
    with open(dst, "w", encoding="utf-8") as f:
        f.write(text)
    rec = {"path": p, "hash": h, "method": method if text.strip() else "error",
           "chars": len(re.sub(r"\s", "", text)), "err": err}
    with LOCK, open(LOG, "a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")


def main():
    os.makedirs(TXT, exist_ok=True)
    ymin = int(os.environ.get("YEAR_MIN", 0))

    def year(p):
        m = re.search(r"PV du C(?:E|SE) (\d{4})", p)
        return int(m.group(1)) if m else 0

    files = [os.path.join(r, n) for r, _, fs in os.walk(SRC) for n in fs if os.path.splitext(n)[1].lower() in EXTS]
    files = [f for f in files if not ymin or year(f) >= ymin]
    todo = [f for f in files if not os.path.exists(os.path.join(TXT, hashlib.md5(f.encode("utf-8")).hexdigest() + ".txt"))]
    print(f"{len(files)} fichiers dans la source, {len(todo)} à indexer", flush=True)
    with ThreadPoolExecutor(max_workers=int(sys.argv[1]) if len(sys.argv) > 1 else 4) as ex:
        for i, _ in enumerate(ex.map(process, todo), 1):
            if i % 25 == 0:
                print(f"{i}/{len(todo)}", flush=True)
    print("TERMINE", flush=True)


if __name__ == "__main__":
    main()
